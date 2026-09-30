import csv
import datetime
import html
import io
from collections import Counter, defaultdict

import frappe
from frappe import _
from frappe.utils import add_days, cstr, formatdate, get_datetime, getdate, today
from frappe.utils.pdf import get_pdf

try:
	from hrms.hr.utils import get_holiday_list_for_employee
except ImportError:
	get_holiday_list_for_employee = None

MAX_DAYS = 93
PDF_DETAIL_LIMIT = 4000

CATEGORIES = [
	"Present",
	"Present (Not Marked)",
	"Half Day",
	"Work From Home",
	"Clocked In Only",
	"Clocked Out Only",
	"On Leave",
	"Absent",
	"Holiday / Week Off",
]

DETAIL_HEADERS = [
	"Employee", "Employee Name", "Department", "Designation", "Date", "Shift",
	"Status", "Clock In", "Clock Out", "Hours Worked", "Marked Attendance", "Leave Type",
]
STATUS_IDX = 6
DATE_IDX = 4


@frappe.whitelist()
def export_ta_data(start_date, end_date, company, file_format="Excel"):
	for dt in ("Attendance", "Employee Checkin", "Employee"):
		if not frappe.has_permission(dt, "read"):
			frappe.throw(_("You do not have permission to read {0}").format(dt), frappe.PermissionError)

	start = getdate(start_date)
	end = min(getdate(end_date), getdate(today()))
	if start > end:
		frappe.throw(_("From date cannot be after To date"))
	if (end - start).days + 1 > MAX_DAYS:
		frappe.throw(_("Please select a range of {0} days or fewer").format(MAX_DAYS))

	details, summary = build_data(start, end, company)
	base = f"Attendance_{company}_{start}_to_{end}".replace(" ", "_")
	title = f"Attendance: {company} ({formatdate(start)} to {formatdate(end)})"

	if file_format == "CSV":
		content, ext = to_csv(details), "csv"
	elif file_format == "PDF":
		content, ext = to_pdf(title, details, summary), "pdf"
	else:
		content, ext = to_xlsx(details, summary), "xlsx"

	frappe.response["filename"] = f"{base}.{ext}"
	frappe.response["filecontent"] = content
	frappe.response["type"] = "download"


def build_data(start, end, company):
	employees = [
		e
		for e in frappe.get_all(
			"Employee",
			filters={"company": company},
			fields=[
				"name", "employee_name", "department", "designation", "date_of_joining",
				"relieving_date", "holiday_list", "default_shift", "status",
			],
			order_by="name asc",
		)
		if e.status == "Active" or (e.relieving_date and getdate(e.relieving_date) >= start)
	]
	if not employees:
		frappe.throw(_("No employees found for {0}").format(company))
	emp_ids = [e.name for e in employees]

	attendance = {}
	for a in frappe.get_all(
		"Attendance",
		filters={"company": company, "docstatus": 1, "attendance_date": ["between", [start, end]]},
		fields=["employee", "attendance_date", "status", "shift", "in_time", "out_time", "working_hours", "leave_type"],
	):
		attendance[(a.employee, getdate(a.attendance_date))] = a

	logs = defaultdict(list)
	for c in frappe.get_all(
		"Employee Checkin",
		filters={
			"employee": ["in", emp_ids],
			"time": ["between", [f"{start} 00:00:00", f"{add_days(end, 1)} 23:59:59"]],
		},
		fields=["employee", "time", "log_type", "shift", "shift_start"],
		order_by="time asc",
	):
		day = getdate(c.shift_start) if c.shift_start else getdate(c.time)
		if start <= day <= end:
			logs[(c.employee, day)].append(c)

	company_holiday_list = frappe.get_cached_value("Company", company, "default_holiday_list")
	holiday_cache = {}

	def holidays_for(emp):
		hl = None
		if get_holiday_list_for_employee:
			try:
				hl = get_holiday_list_for_employee(emp.name, raise_exception=False)
			except Exception:
				hl = None
		hl = hl or emp.holiday_list or company_holiday_list
		if not hl:
			return set()
		if hl not in holiday_cache:
			holiday_cache[hl] = {
				getdate(d)
				for d in frappe.get_all(
					"Holiday",
					filters={"parent": hl, "holiday_date": ["between", [start, end]]},
					pluck="holiday_date",
				)
			}
		return holiday_cache[hl]

	dates, d = [], start
	while d <= end:
		dates.append(getdate(d))
		d = add_days(d, 1)

	details = []
	summary = {d: Counter() for d in dates}

	for emp in employees:
		hol = holidays_for(emp)
		doj = getdate(emp.date_of_joining) if emp.date_of_joining else None
		rel = getdate(emp.relieving_date) if emp.relieving_date else None

		for d in dates:
			if (doj and d < doj) or (rel and d > rel):
				continue

			a = attendance.get((emp.name, d))
			lg = logs.get((emp.name, d), [])
			ins = [l.time for l in lg if l.log_type == "IN"]
			outs = [l.time for l in lg if l.log_type == "OUT"]
			untyped = [l.time for l in lg if not l.log_type]

			clock_in = ins[0] if ins else (untyped[0] if untyped else None)
			clock_out = outs[-1] if outs else (untyped[-1] if len(untyped) > 1 else None)
			if a:
				clock_in = a.in_time or clock_in
				clock_out = a.out_time or clock_out

			status = resolve_status(a, clock_in, clock_out, d in hol)

			hours = ""
			if clock_in and clock_out and get_datetime(clock_out) > get_datetime(clock_in):
				hours = round((get_datetime(clock_out) - get_datetime(clock_in)).total_seconds() / 3600, 2)
			elif a and a.working_hours:
				hours = round(a.working_hours, 2)

			shift = (a and a.shift) or next((l.shift for l in lg if l.shift), None) or emp.default_shift or ""

			details.append([
				emp.name, emp.employee_name, emp.department or "", emp.designation or "", d, shift,
				status, fmt_time(clock_in, d), fmt_time(clock_out, d), hours,
				a.status if a else "Not Marked", (a and a.leave_type) or "",
			])
			summary[d][status] += 1

	details.sort(key=lambda r: (r[DATE_IDX], r[0]))
	return details, summary


def resolve_status(a, clock_in, clock_out, is_holiday):
	if a:
		if a.status in ("Present", "Half Day", "Work From Home") and clock_in and not clock_out:
			return "Clocked In Only"
		return a.status
	if clock_in and clock_out:
		return "Present (Not Marked)"
	if clock_in:
		return "Clocked In Only"
	if clock_out:
		return "Clocked Out Only"
	if is_holiday:
		return "Holiday / Week Off"
	return "Absent"


def fmt_time(value, day):
	if not value:
		return ""
	dt = get_datetime(value)
	return dt.strftime("%H:%M") if dt.date() == day else dt.strftime("%d-%m-%Y %H:%M")


def summary_table(summary):
	extra = sorted({c for cnt in summary.values() for c in cnt} - set(CATEGORIES))
	cats = CATEGORIES + extra
	headers = ["Date", "Headcount"] + cats
	rows, totals = [], Counter()
	for d in sorted(summary):
		cnt = summary[d]
		rows.append([d, sum(cnt.values())] + [cnt.get(c, 0) for c in cats])
		totals.update(cnt)
	rows.append(["Total", sum(totals.values())] + [totals.get(c, 0) for c in cats])
	return headers, rows


def to_xlsx(details, summary):
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill
	from openpyxl.utils import get_column_letter

	wb = Workbook()
	wb.remove(wb.active)

	def add_sheet(title, headers, rows, date_col):
		ws = wb.create_sheet(title)
		ws.append(headers)
		for cell in ws[1]:
			cell.font = Font(bold=True, color="FFFFFF")
			cell.fill = PatternFill("solid", fgColor="1F4E78")
		for r in rows:
			ws.append(r)
		ws.freeze_panes = "A2"
		ws.auto_filter.ref = ws.dimensions
		for row in ws.iter_rows(min_row=2, min_col=date_col, max_col=date_col):
			for cell in row:
				if isinstance(cell.value, datetime.date):
					cell.number_format = "DD-MM-YYYY"
		for i, h in enumerate(headers, 1):
			width = max([len(cstr(h))] + [len(cstr(r[i - 1])) for r in rows[:1000]]) + 2
			ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 40)

	s_headers, s_rows = summary_table(summary)
	add_sheet("Summary", s_headers, s_rows, 1)
	add_sheet("All Records", DETAIL_HEADERS, details, DATE_IDX + 1)

	groups = {
		"Present": ("Present", "Present (Not Marked)", "Half Day", "Work From Home"),
		"Clocked In Only": ("Clocked In Only",),
		"Clocked Out Only": ("Clocked Out Only",),
		"Absent": ("Absent",),
		"On Leave": ("On Leave",),
	}
	for sheet, statuses in groups.items():
		add_sheet(sheet, DETAIL_HEADERS, [r for r in details if r[STATUS_IDX] in statuses], DATE_IDX + 1)

	out = io.BytesIO()
	wb.save(out)
	return out.getvalue()


def to_csv(details):
	buf = io.StringIO()
	w = csv.writer(buf)
	w.writerow(DETAIL_HEADERS)
	for r in details:
		w.writerow([formatdate(c) if isinstance(c, datetime.date) else c for c in r])
	return buf.getvalue().encode("utf-8-sig")


def to_pdf(title, details, summary):
	def cell(v):
		if isinstance(v, datetime.date):
			v = formatdate(v)
		return html.escape(cstr(v))

	def table(headers,
